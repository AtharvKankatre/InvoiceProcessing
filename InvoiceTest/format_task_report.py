import openpyxl
from openpyxl.styles import PatternFill, Font

path = 'd:\\Inpinite\\InvoiceTest\\Cooper_Task_Report.xlsx'
wb = openpyxl.load_workbook(path)
ws = wb.active

# 1. Remove the "Period" row
# Iterate backwards to safely delete rows
for i in range(ws.max_row, 0, -1):
    cell_val = str(ws.cell(row=i, column=1).value).strip() if ws.cell(row=i, column=1).value else ""
    if cell_val == "Period":
        ws.delete_rows(i)

# 2. Color "Completed" status column green
# Excel's standard "Good" style: Light green fill with dark green text
green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
green_font = Font(color="006100")

# Standard "Pending" style (optional, but good for completeness)
# yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
# yellow_font = Font(color="9C5700")

for row in ws.iter_rows(min_row=2):
    # In my CSV, Status is usually the 4th column, but some rows might differ.
    # Let's check all cells in the row just to be safe, or just check the last few.
    for cell in row:
        if str(cell.value).strip() == "Completed":
            cell.fill = green_fill
            cell.font = green_font
        # elif str(cell.value).strip() == "Pending":
        #    cell.fill = yellow_fill
        #    cell.font = yellow_font

wb.save(path)
print("Formatting applied successfully.")
