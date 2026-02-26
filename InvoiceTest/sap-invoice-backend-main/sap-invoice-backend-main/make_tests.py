import os
import pandas as pd
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

TEMPLATE_COLUMNS = ['part_number', 'date', 'qty', 'usd_rate', 'conversion_rate', 'retail_invoice_number']
COLUMN_DESCRIPTIONS = [
    'Retail Part Number (e.g. RETAIL-X100)',
    'Date (YYYY-MM-DD)',
    'Quantity to dispatch',
    'Selling USD rate per unit',
    'INR/USD Conversion (FX) rate',
    'Retail Invoice Number (must be unique)'
]

def make_wb(data):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'FX Bulk Upload'

    header_font = Font(bold=True, color='FFFFFF', size=12)
    header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
    desc_font = Font(italic=True, color='808080', size=10)
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for col_idx, header in enumerate(TEMPLATE_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border

    for col_idx, desc in enumerate(COLUMN_DESCRIPTIONS, start=1):
        cell = ws.cell(row=2, column=col_idx, value=desc)
        cell.font = desc_font
        cell.alignment = Alignment(horizontal='center')

    widths = [25, 15, 12, 15, 18, 28]
    for col_idx, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = w

    for i, row in enumerate(data, start=3):
        for j, key in enumerate(TEMPLATE_COLUMNS, start=1):
            ws.cell(row=i, column=j, value=row[key])
    
    return wb

time_str = datetime.now().strftime('%H%M%S')

# File 1: Single Row
t1 = [{
    'part_number': 'PART-X100',
    'date': '2025-02-01',
    'qty': 5,
    'usd_rate': 100.50,
    'conversion_rate': 83.2,
    'retail_invoice_number': f'RET-TEST-{time_str}-1'
}]
make_wb(t1).save('test_single_row.xlsx')

# File 2: Multi Row
t2 = [
    {
        'part_number': 'PART-X100',
        'date': '2025-02-02',
        'qty': 2,
        'usd_rate': 100.50,
        'conversion_rate': 83.2,
        'retail_invoice_number': f'RET-TEST-{time_str}-2'
    },
    {
        'part_number': 'PART-X100',
        'date': '2025-02-03',
        'qty': 3,
        'usd_rate': 105.00,
        'conversion_rate': 84.0,
        'retail_invoice_number': f'RET-TEST-{time_str}-3'
    }
]
make_wb(t2).save('test_multi_row.xlsx')

print('Generated test_single_row.xlsx and test_multi_row.xlsx')
