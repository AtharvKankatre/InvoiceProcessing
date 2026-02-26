# coding: utf-8
"""Test the enhanced Consumption Details export."""
import os, json, django
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'sap_invoice.settings')
django.setup()

from urllib.request import Request, urlopen
from urllib.error import HTTPError
from django.contrib.auth import get_user_model
from rest_framework_simplejwt.tokens import RefreshToken

user = get_user_model().objects.first()
token = str(RefreshToken.for_user(user).access_token)

# Call export endpoint
url = "http://localhost:8000/api/sales/export-stock-ledger/"
req = Request(url, headers={"Authorization": "Bearer " + token})

try:
    resp = urlopen(req, timeout=30)
    content = resp.read()
    
    # Save to file
    out_path = r"D:\Inpinite\InvoiceTest\test_export_output.xlsx"
    with open(out_path, "wb") as f:
        f.write(content)
    
    # Verify with openpyxl
    import openpyxl
    wb = openpyxl.load_workbook(out_path)
    print(f"Sheets: {wb.sheetnames}")
    
    if "Consumption Details" in wb.sheetnames:
        ws = wb["Consumption Details"]
        print(f"\nConsumption Details: {ws.max_row} rows x {ws.max_column} columns")
        print(f"\nHeaders (row 1):")
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col)
            print(f"  Col {col}: {cell.value}")
        
        print(f"\nFirst 20 data rows:")
        for row in range(2, min(22, ws.max_row + 1)):
            vals = []
            row_type = ""
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=row, column=col)
                v = cell.value
                if v is not None:
                    vals.append(str(v)[:20])
                else:
                    vals.append("")
                # Detect row type from fill
                if cell.fill and cell.fill.start_color and cell.fill.start_color.rgb:
                    rgb = str(cell.fill.start_color.rgb)
                    if "D6E4F0" in rgb:
                        row_type = "HEADER"
                    elif "E2EFDA" in rgb:
                        row_type = "SUBTOTAL"
                    elif "FFC000" in rgb:
                        row_type = "GRAND"
            print(f"  Row {row:3d} [{row_type:8s}]: {' | '.join(vals[:7])}")
        
        print(f"\nLast 3 rows:")
        for row in range(max(2, ws.max_row - 2), ws.max_row + 1):
            vals = []
            for col in range(1, min(8, ws.max_column + 1)):
                v = ws.cell(row=row, column=col).value
                vals.append(str(v)[:20] if v else "")
            print(f"  Row {row:3d}: {' | '.join(vals)}")
    
    if "Stock Ledger" in wb.sheetnames:
        ws2 = wb["Stock Ledger"]
        print(f"\nStock Ledger: {ws2.max_row} rows x {ws2.max_column} columns")
    
    print(f"\nSaved to: {out_path}")
    print("SUCCESS!")
    
except HTTPError as e:
    body = e.read().decode()
    print(f"ERROR {e.code}: {body[:500]}")
except Exception as e:
    print(f"ERROR: {e}")
    import traceback
    traceback.print_exc()
