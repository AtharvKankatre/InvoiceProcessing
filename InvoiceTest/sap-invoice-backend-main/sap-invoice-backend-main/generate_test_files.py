# coding: utf-8
"""
Step 1: Query the database for existing data to create realistic test files.
Step 2: Generate 10 bulk upload test files for different scenarios.
"""
import os, sys, json, django
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'sap_invoice.settings')
django.setup()

from sales.models import Invoice, InvoiceRetailPartMap
from retail.models import InvoiceEntry
from decimal import Decimal

OUT_DIR = r"D:\Inpinite\InvoiceTest\sap-invoice-backend-main\sap-invoice-backend-main\TestingFiles"
os.makedirs(OUT_DIR, exist_ok=True)

print("=" * 60)
print("STEP 1: DATABASE INVENTORY")
print("=" * 60)

# Get real part numbers with stock
parts = InvoiceRetailPartMap.objects.values_list('retail_part_number', 'sale_part_number', 'company_name').distinct()
print(f"\nPart Mappings ({len(parts)}):")
part_list = []
for rp, sp, co in parts:
    # Calculate available stock: Invoice entries where part_number == sale_part_number
    total_in = Invoice.objects.filter(part_number=sp).values_list('qty', flat=True)
    in_qty = sum(q for q in total_in if q)
    
    total_out = InvoiceEntry.objects.filter(part_number=rp).values_list('qty', flat=True)
    out_qty = sum(q for q in total_out if q)
    
    avail = in_qty - out_qty
    print(f"  {rp:25s} (SAP: {sp:15s}) | Company: {co:20s} | Stock: {avail}")
    part_list.append({'retail': rp, 'sale': sp, 'company': co, 'available': avail})

# Get existing retail invoice numbers
existing_inv = set(x for x in InvoiceEntry.objects.values_list('retail_invoice_number', flat=True).distinct() if x)
print(f"\nExisting retail invoice numbers: {len(existing_inv)}")
for inv in sorted(existing_inv)[:10]:
    print(f"  {inv}")
if len(existing_inv) > 10:
    print(f"  ... and {len(existing_inv) - 10} more")

# Get customers
customers = Invoice.objects.values('customer_name', 'customer_code').distinct()
print(f"\nCustomers ({len(customers)}):")
for c in customers:
    print(f"  {c['customer_code'] or 'N/A':15s} | {c['customer_name']}")

# Save inventory data for the test generator
inventory_data = {
    'parts': part_list,
    'existing_invoices': list(existing_inv)[:20],
    'customer_count': len(customers),
}

inv_path = os.path.join(OUT_DIR, '_inventory_snapshot.json')
with open(inv_path, 'w') as f:
    json.dump(inventory_data, f, indent=2, default=str)
print(f"\nSaved inventory snapshot to: {inv_path}")

# ═══════════════════════════════════════════════════════════
# STEP 2: Generate 10 Bulk Upload Test Files
# ═══════════════════════════════════════════════════════════
print("\n" + "=" * 60)
print("STEP 2: GENERATING 10 BULK UPLOAD TEST FILES")
print("=" * 60)

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import date, timedelta
import random
import time

ts = str(int(time.time()))[-6:]

def create_workbook(rows_data, filename, description):
    """Create a styled Excel file for bulk upload testing."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bulk Upload"
    
    headers = ["part_number", "date", "qty", "usd_rate", "conversion_rate", "retail_invoice_number"]
    hfill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    hfont = Font(color="FFFFFF", bold=True, size=11)
    
    for i, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.fill = hfill
        cell.font = hfont
        cell.alignment = Alignment(horizontal="center")
    
    for r_idx, row in enumerate(rows_data, 2):
        for c_idx, val in enumerate(row, 1):
            ws.cell(row=r_idx, column=c_idx, value=val)
    
    # Column widths
    for i, w in enumerate([22, 14, 10, 12, 16, 28], 1):
        ws.column_dimensions[chr(64 + i)].width = w
    
    # Add a description sheet
    desc_ws = wb.create_sheet("Test Info")
    desc_ws.cell(row=1, column=1, value="Test File Description").font = Font(bold=True, size=14)
    desc_ws.cell(row=2, column=1, value=description)
    desc_ws.cell(row=3, column=1, value=f"Generated: {date.today().strftime('%d-%m-%Y')}")
    desc_ws.column_dimensions['A'].width = 80
    
    filepath = os.path.join(OUT_DIR, filename)
    wb.save(filepath)
    print(f"  Created: {filename} ({len(rows_data)} rows) — {description}")
    return filepath

# Use real parts if available, else use placeholders
if part_list:
    # Pick parts that have stock
    parts_with_stock = [p for p in part_list if p['available'] > 0]
    if not parts_with_stock:
        parts_with_stock = part_list[:3]
    p1 = parts_with_stock[0]['retail'] if len(parts_with_stock) > 0 else "RETAIL-A400"
    p2 = parts_with_stock[1]['retail'] if len(parts_with_stock) > 1 else "RETAIL-X100"
    p3 = parts_with_stock[2]['retail'] if len(parts_with_stock) > 2 else "RETAIL-Y200"
    max_stock_1 = parts_with_stock[0]['available'] if len(parts_with_stock) > 0 else 100
    max_stock_2 = parts_with_stock[1]['available'] if len(parts_with_stock) > 1 else 50
else:
    p1, p2, p3 = "RETAIL-A400", "RETAIL-X100", "RETAIL-Y200"
    max_stock_1, max_stock_2 = 100, 50

d = date(2026, 2, 24)

# ── FILE 1: Happy Path (3 valid rows, all should pass) ──
create_workbook([
    (p1, d, min(10, max_stock_1), 10.5, 83.2, f"HT-{ts}-001"),
    (p2, d, min(5, max_stock_2), 12.0, 83.5, f"HT-{ts}-002"),
    (p1, d, min(10, max_stock_1), 11.0, 83.0, f"HT-{ts}-003"),
], "01_happy_path.xlsx", "All 3 rows are valid. Should preview with 0 errors.")

# ── FILE 2: Missing Fields (required fields left blank) ──
create_workbook([
    (p1, d, 10, 10.5, 83.2, f"MF-{ts}-001"),     # Valid
    (None, d, 10, 10.5, 83.2, f"MF-{ts}-002"),    # Missing part_number
    (p1, None, 10, 10.5, 83.2, f"MF-{ts}-003"),   # Missing date
    (p1, d, None, 10.5, 83.2, f"MF-{ts}-004"),    # Missing qty
    (p1, d, 10, None, 83.2, f"MF-{ts}-005"),      # Missing usd_rate
    (p1, d, 10, 10.5, None, f"MF-{ts}-006"),      # Missing conversion_rate
    (p1, d, 10, 10.5, 83.2, None),                 # Missing retail_invoice_number
], "02_missing_fields.xlsx", "Tests each required field missing one at a time. Rows 2-7 should fail.")

# ── FILE 3: Invalid Part Numbers ──
create_workbook([
    (p1, d, 5, 10.5, 83.2, f"IP-{ts}-001"),                # Valid
    ("FAKE-PART-999", d, 5, 10.5, 83.2, f"IP-{ts}-002"),   # Fake part
    ("NONEXISTENT", d, 5, 10.5, 83.2, f"IP-{ts}-003"),     # Non-existent
    ("", d, 5, 10.5, 83.2, f"IP-{ts}-004"),                # Empty string
], "03_invalid_parts.xlsx", "Tests invalid/fake part numbers. Rows 2-4 should fail with 'Part not found'.")

# ── FILE 4: Duplicate Invoice Numbers (within batch) ──
create_workbook([
    (p1, d, 5, 10.5, 83.2, f"DI-{ts}-001"),
    (p2, d, 5, 12.0, 83.5, f"DI-{ts}-001"),   # Same invoice as row 1
    (p1, d, 5, 10.5, 83.2, f"DI-{ts}-002"),
    (p2, d, 5, 12.0, 83.5, f"DI-{ts}-002"),   # Same invoice as row 3
], "04_duplicate_invoices.xlsx", "Tests duplicate retail_invoice_number within same file. Rows 2 & 4 should fail.")

# ── FILE 5: Stock Exceeded (qty > available stock) ──
create_workbook([
    (p1, d, 999999, 10.5, 83.2, f"SE-{ts}-001"),   # Way too much
    (p1, d, min(5, max_stock_1), 10.5, 83.2, f"SE-{ts}-002"),   # Small valid qty
], "05_stock_exceeded.xlsx", "Tests qty > available stock. Row 1 should fail with stock error.")

# ── FILE 6: Negative and Zero Values ──
create_workbook([
    (p1, d, -10, 10.5, 83.2, f"NZ-{ts}-001"),     # Negative qty
    (p1, d, 0, 10.5, 83.2, f"NZ-{ts}-002"),       # Zero qty
    (p1, d, 5, -5.0, 83.2, f"NZ-{ts}-003"),       # Negative usd_rate
    (p1, d, 5, 0, 83.2, f"NZ-{ts}-004"),          # Zero usd_rate
    (p1, d, 5, 10.5, -83.2, f"NZ-{ts}-005"),      # Negative conversion_rate
    (p1, d, 5, 10.5, 0, f"NZ-{ts}-006"),          # Zero conversion_rate
], "06_negative_zero_values.xlsx", "Tests negative and zero values for qty, usd_rate, conversion_rate. All should fail.")

# ── FILE 7: Fractional Quantities ──
create_workbook([
    (p1, d, 10.5, 10.5, 83.2, f"FQ-{ts}-001"),    # 10.5 qty — not integer
    (p1, d, 3.7, 10.5, 83.2, f"FQ-{ts}-002"),     # 3.7 qty — not integer
    (p1, d, 10.0, 10.5, 83.2, f"FQ-{ts}-003"),    # 10.0 — integer disguised as float
    (p1, d, 5, 10.5, 83.2, f"FQ-{ts}-004"),       # Valid integer
], "07_fractional_qty.xlsx", "Tests fractional quantities. Rows 1-2 should fail (not whole number). Rows 3-4 should pass.")

# ── FILE 8: Large Batch (50 rows) ──
large_rows = []
for i in range(50):
    part = p1 if i % 2 == 0 else p2
    qty = 1  # Minimal qty to avoid stock issues
    large_rows.append((part, d, qty, 10.5 + (i * 0.1), 83.0 + (i * 0.01), f"LB-{ts}-{i+1:03d}"))
create_workbook(large_rows, "08_large_batch_50.xlsx", "50-row batch to test performance. All rows should be valid (qty=1 each).")

# ── FILE 9: Mixed Valid and Invalid ──
create_workbook([
    (p1, d, min(5, max_stock_1), 10.5, 83.2, f"MX-{ts}-001"),          # Valid
    ("FAKE-PART", d, 5, 10.5, 83.2, f"MX-{ts}-002"),                   # Bad part
    (p2, d, min(3, max_stock_2), 12.0, 83.5, f"MX-{ts}-003"),          # Valid
    (p1, None, 5, 10.5, 83.2, f"MX-{ts}-004"),                         # Missing date
    (p1, d, min(2, max_stock_1), 11.0, 83.0, f"MX-{ts}-005"),          # Valid
    (p1, d, -5, 10.5, 83.2, f"MX-{ts}-006"),                           # Negative qty
    (p2, d, min(2, max_stock_2), 12.5, 83.1, f"MX-{ts}-007"),          # Valid
    (p1, d, 5, 10.5, 83.2, f"MX-{ts}-001"),                            # Dup invoice with row 1
], "09_mixed_valid_invalid.xlsx", "Mix of valid and invalid rows. Rows 1,3,5,7 valid; rows 2,4,6,8 invalid.")

# ── FILE 10: DB Duplicate Invoice (uses an existing invoice number from DB) ──
existing_list = list(existing_inv)
db_dup_rows = [(p1, d, min(5, max_stock_1), 10.5, 83.2, f"DD-{ts}-001")]  # Valid
if existing_list:
    db_dup_rows.append((p1, d, min(3, max_stock_1), 10.5, 83.2, existing_list[0]))  # DB duplicate
    desc10 = f"Row 2 uses existing invoice '{existing_list[0]}' from DB. Should fail with 'Already exists'."
else:
    db_dup_rows.append((p1, d, min(3, max_stock_1), 10.5, 83.2, "BULK-TEST-001"))
    desc10 = "Row 2 uses potentially existing invoice. May fail with 'Already exists'."
create_workbook(db_dup_rows, "10_db_duplicate_invoice.xlsx", desc10)

# ═══════════════════════════════════════════════════════════
# Summary
# ═══════════════════════════════════════════════════════════
print(f"\n{'=' * 60}")
print(f"Generated 10 test files in: {OUT_DIR}")
print(f"Timestamp prefix: {ts}")
print(f"Parts used: {p1}, {p2}")
print(f"{'=' * 60}")
