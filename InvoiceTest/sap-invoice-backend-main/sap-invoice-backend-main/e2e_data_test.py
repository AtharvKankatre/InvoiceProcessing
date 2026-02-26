# coding: utf-8
"""
E2E Data Flow Verification — Tests all features:
1. Bulk Upload Preview (all 10 test files)
2. Stock Ledger Export (both sheets)
3. Part History Export
"""
import os, sys, json, django, glob
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'sap_invoice.settings')
django.setup()

from urllib.request import Request, urlopen
from urllib.error import HTTPError
from django.contrib.auth import get_user_model
from rest_framework_simplejwt.tokens import RefreshToken

BASE = "http://localhost:8000"
TEST_DIR = r"D:\Inpinite\InvoiceTest\sap-invoice-backend-main\sap-invoice-backend-main\TestingFiles"
PASS = 0
FAIL = 0

user = get_user_model().objects.first()
token = str(RefreshToken.for_user(user).access_token)

def test(name, cond, detail=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  PASS: {name}")
    else:
        FAIL += 1
        print(f"  FAIL: {name} -- {detail}")

def api_post_file(path, filepath):
    bd = "----Boundary"
    fn = os.path.basename(filepath)
    with open(filepath, "rb") as f:
        fd = f.read()
    body_str = "--" + bd + "\r\nContent-Disposition: form-data; name=\"file\"; filename=\"" + fn + "\"\r\nContent-Type: application/octet-stream\r\n\r\n"
    body = body_str.encode() + fd + ("\r\n--" + bd + "--\r\n").encode()
    req = Request(f"{BASE}{path}", data=body, headers={
        "Content-Type": "multipart/form-data; boundary=" + bd,
        "Authorization": "Bearer " + token
    }, method="POST")
    try:
        resp = urlopen(req, timeout=30)
        return resp.status, json.loads(resp.read().decode())
    except HTTPError as e:
        body = e.read()
        try:
            return e.code, json.loads(body.decode())
        except:
            return e.code, {"error": body.decode()[:200]}

def api_get(path):
    req = Request(f"{BASE}{path}", headers={"Authorization": "Bearer " + token})
    try:
        resp = urlopen(req, timeout=30)
        return resp.status, resp.read()
    except HTTPError as e:
        return e.code, e.read()

# ═══════════════════════════════════════════════════════════
# TEST 1: BULK UPLOAD PREVIEW (10 files)
# ═══════════════════════════════════════════════════════════
print("=" * 70)
print("TEST 1: BULK UPLOAD PREVIEW")
print("=" * 70)

expected = {
    "01_happy_path.xlsx":           {"valid_min": 2, "invalid_max": 1},
    "02_missing_fields.xlsx":       {"valid_min": 1, "invalid_min": 5},
    "03_invalid_parts.xlsx":        {"valid_min": 1, "invalid_min": 2},
    "04_duplicate_invoices.xlsx":   {"valid_min": 2, "invalid_min": 2},
    "05_stock_exceeded.xlsx":       {"invalid_min": 1},
    "06_negative_zero_values.xlsx": {"invalid_min": 2},
    "07_fractional_qty.xlsx":       {"valid_min": 1},
    "08_large_batch_50.xlsx":       {"total": 50},
    "09_mixed_valid_invalid.xlsx":  {"valid_min": 3, "invalid_min": 3},
    "10_db_duplicate_invoice.xlsx": {"valid_min": 1},
}

for filename in sorted(glob.glob(os.path.join(TEST_DIR, "[0-9]*.xlsx"))):
    base = os.path.basename(filename)
    
    print(f"\n  --- {base} ---")
    code, data = api_post_file("/api/retail/bulk-preview/", filename)
    test(f"{base}: HTTP 200", code == 200, f"Got {code}")
    
    if code != 200:
        print(f"    Error: {json.dumps(data)[:200]}")
        continue
    
    rows = data.get("rows", [])
    valid = [r for r in rows if r.get("is_valid")]
    invalid = [r for r in rows if not r.get("is_valid")]
    
    print(f"    Rows: {len(rows)} total, {len(valid)} valid, {len(invalid)} invalid")
    
    exp = expected.get(base, {})
    if "total" in exp:
        test(f"{base}: has {exp['total']} rows", len(rows) == exp['total'], f"Got {len(rows)}")
    if "valid_min" in exp:
        test(f"{base}: >= {exp['valid_min']} valid", len(valid) >= exp['valid_min'], f"Got {len(valid)}")
    if "invalid_min" in exp:
        test(f"{base}: >= {exp['invalid_min']} invalid", len(invalid) >= exp['invalid_min'], f"Got {len(invalid)}")
    if "invalid_max" in exp:
        test(f"{base}: <= {exp['invalid_max']} invalid", len(invalid) <= exp['invalid_max'], f"Got {len(invalid)}")
    
    # Show errors for invalid rows
    for r in invalid[:3]:
        errs = r.get("errors", {})
        print(f"    Row {r['row_id']}: {json.dumps(errs)[:120]}")

# ═══════════════════════════════════════════════════════════
# TEST 2: STOCK LEDGER EXPORT
# ═══════════════════════════════════════════════════════════
print("\n" + "=" * 70)
print("TEST 2: STOCK LEDGER EXPORT")
print("=" * 70)

code, content = api_get("/api/sales/export-stock-ledger/")
test("Stock Ledger export: HTTP 200", code == 200, f"Got {code}")

if code == 200:
    out = os.path.join(TEST_DIR, "export_stock_ledger.xlsx")
    with open(out, "wb") as f:
        f.write(content)
    
    import openpyxl
    wb = openpyxl.load_workbook(out)
    test("Has 'Stock Ledger' sheet", "Stock Ledger" in wb.sheetnames)
    test("Has 'Consumption Details' sheet", "Consumption Details" in wb.sheetnames)
    
    if "Stock Ledger" in wb.sheetnames:
        ws1 = wb["Stock Ledger"]
        test("Stock Ledger has data (>1 row)", ws1.max_row > 1, f"Only {ws1.max_row} rows")
        test("Stock Ledger has 14 columns", ws1.max_column == 14, f"Got {ws1.max_column} cols")
        
        # Check headers
        h1 = ws1.cell(row=1, column=1).value
        h2 = ws1.cell(row=1, column=2).value
        test("Header col 1 = 'Code No. Stock AC'", h1 == "Code No. Stock AC", f"Got '{h1}'")
        test("Header col 2 = 'Party Name'", h2 == "Party Name", f"Got '{h2}'")
        print(f"    Stock Ledger: {ws1.max_row} rows x {ws1.max_column} columns")
    
    if "Consumption Details" in wb.sheetnames:
        ws2 = wb["Consumption Details"]
        test("Consumption Details has data (>1 row)", ws2.max_row > 1, f"Only {ws2.max_row} rows")
        test("Consumption Details has 16 columns", ws2.max_column == 16, f"Got {ws2.max_column} cols")
        
        # Check customer grouping
        has_header_row = False
        has_subtotal_row = False
        has_grand_total = False
        for row in range(2, min(ws2.max_row + 1, 100)):
            cell = ws2.cell(row=row, column=2)
            fill = cell.fill
            val = cell.value or ""
            if fill and fill.start_color and "D6E4F0" in str(fill.start_color.rgb):
                has_header_row = True
            if fill and fill.start_color and "E2EFDA" in str(fill.start_color.rgb):
                has_subtotal_row = True
            if "GRAND TOTAL" in str(val):
                has_grand_total = True
        
        test("Has customer header rows (blue)", has_header_row)
        test("Has subtotal rows (green)", has_subtotal_row)
        test("Has GRAND TOTAL row", has_grand_total)
        
        # Check new columns exist
        h8 = ws2.cell(row=1, column=8).value
        h11 = ws2.cell(row=1, column=11).value
        h14 = ws2.cell(row=1, column=14).value
        test("Column 8 = 'Opening Qty'", h8 == "Opening Qty", f"Got '{h8}'")
        test("Column 11 = 'Shipment Qty'", h11 == "Shipment Qty", f"Got '{h11}'")
        test("Column 14 = 'Closing Qty'", h14 == "Closing Qty", f"Got '{h14}'")
        
        print(f"    Consumption Details: {ws2.max_row} rows x {ws2.max_column} columns")
    
    print(f"    Saved: {out}")

# ═══════════════════════════════════════════════════════════
# TEST 3: PART HISTORY EXPORT
# ═══════════════════════════════════════════════════════════
print("\n" + "=" * 70)
print("TEST 3: PART HISTORY EXPORT (All Parts)")
print("=" * 70)

# First find a real part number
from sales.models import InvoiceRetailPartMap
first_map = InvoiceRetailPartMap.objects.first()
if first_map:
    part = first_map.retail_part_number
    code, content = api_get(f"/api/sales/parts/{part}/export/?scope=all")
    test("Part History ALL export: HTTP 200", code == 200, f"Got {code}")
    
    if code == 200:
        out = os.path.join(TEST_DIR, "export_part_history_all.xlsx")
        with open(out, "wb") as f:
            f.write(content)
        
        wb = openpyxl.load_workbook(out)
        ws = wb.active
        test("Part History has data", ws.max_row > 1, f"Only {ws.max_row} rows")
        
        # Check for new columns (FC Value, INR Value, Exchange Rate)
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        test("Has 'FC Value' column", "FC Value" in headers, f"Headers: {headers}")
        test("Has 'INR Value' column", "INR Value" in headers, f"Headers: {headers}")
        test("Has 'Exchange Rate' column", "Exchange Rate" in headers, f"Headers: {headers}")
        
        print(f"    Part History: {ws.max_row} rows x {ws.max_column} columns")
        print(f"    Headers: {headers}")
        print(f"    Saved: {out}")
else:
    print("  SKIP: No part mappings in database")

# ═══════════════════════════════════════════════════════════
# SUMMARY
# ═══════════════════════════════════════════════════════════
print("\n" + "=" * 70)
total = PASS + FAIL
print(f"RESULTS: {PASS}/{total} passed, {FAIL} failed")
if FAIL == 0:
    print("ALL TESTS PASSED!")
else:
    print(f"{FAIL} test(s) need attention")
print("=" * 70)
