# coding: utf-8
"""Regenerate test file with unique invoice numbers to avoid DB collision."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import date
import time

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Test Scenarios"

headers = ["part_number", "date", "qty", "usd_rate", "conversion_rate", "retail_invoice_number"]
header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True, size=11)

for col_idx, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_idx, value=header)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center")

# Unique prefix to avoid DB collision
ts = str(int(time.time()))[-6:]
test_rows = [
    ("RETAIL-A400", date(2026, 2, 24), 100, 10.5, 83.2, f"UT-{ts}-001"),
    ("RETAIL-X100", date(2026, 2, 24), 50, 12.0, 83.5, f"UT-{ts}-002"),
    ("RETAIL-A400", date(2026, 2, 24), 200, 10.5, 83.2, f"UT-{ts}-003"),
    ("RETAIL-A400", None, 50, 10.5, 83.2, f"UT-{ts}-004"),
    (None, date(2026, 2, 24), 50, 10.5, 83.2, f"UT-{ts}-005"),
    ("FAKE-PART-999", date(2026, 2, 24), 50, 10.5, 83.2, f"UT-{ts}-006"),
    ("RETAIL-A400", date(2026, 2, 24), None, 10.5, 83.2, f"UT-{ts}-007"),
    ("RETAIL-A400", date(2026, 2, 24), 0, 10.5, 83.2, f"UT-{ts}-008"),
    ("RETAIL-A400", date(2026, 2, 24), 50, None, 83.2, f"UT-{ts}-009"),
    ("RETAIL-A400", date(2026, 2, 24), 50, 10.5, None, f"UT-{ts}-010"),
    ("RETAIL-A400", date(2026, 2, 24), 50, 10.5, 83.2, None),
    ("RETAIL-Y200", date(2026, 2, 24), 50, 11.0, 83.0, f"UT-{ts}-001"),
    ("RETAIL-Y200", date(2026, 2, 24), 50, 11.0, 83.0, "BULK-TEST-001"),
    ("RETAIL-A400", date(2026, 2, 24), 999999, 10.5, 83.2, f"UT-{ts}-014"),
]

for row_idx, row_data in enumerate(test_rows, 2):
    for col_idx, value in enumerate(row_data, 1):
        ws.cell(row=row_idx, column=col_idx, value=value)

ws.column_dimensions['A'].width = 20
ws.column_dimensions['B'].width = 14
ws.column_dimensions['C'].width = 10
ws.column_dimensions['D'].width = 12
ws.column_dimensions['E'].width = 16
ws.column_dimensions['F'].width = 25

output_path = r"D:\Inpinite\InvoiceTest\STAGING_TEST_FILE.xlsx"
wb.save(output_path)
print(f"Created: {output_path} (prefix: UT-{ts})")
